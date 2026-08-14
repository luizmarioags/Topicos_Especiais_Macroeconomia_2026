"""Extrai valores em cache e fórmulas das planilhas originais para validação.

Este utilitário é de desenvolvimento e requer ``openpyxl``. Ele não é
necessário para usar os pacotes R ou Python.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import openpyxl
import pandas as pd


DYNAMIC_COLUMNS = [
    "t",
    "h",
    "y",
    "y_potential",
    "y_autonomous",
    "r_expected",
    "r_cb",
    "r_natural",
    "pi",
    "pi_expected",
    "pi_long_run",
    "pi_target",
    "gamma",
    "theta",
    "alpha",
    "phi_pi",
    "phi_y",
    "chi",
    "eta",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_dynamic(path: Path) -> pd.DataFrame:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["M3E dinâmico"]
    rows = list(
        sheet.iter_rows(
            min_row=18,
            max_row=528,
            min_col=1,
            max_col=19,
            values_only=True,
        )
    )
    return pd.DataFrame(rows, columns=DYNAMIC_COLUMNS)


def array_formula_text(value: object) -> str | None:
    return getattr(value, "text", None)


def extract_static(path: Path) -> dict[str, object]:
    values = openpyxl.load_workbook(path, data_only=True)["M3E numérico"]
    formulas = openpyxl.load_workbook(path, data_only=False)["M3E numérico"]
    return {
        "variable_order": ["y", "pi_expected", "pi", "r_expected"],
        "equilibrium_1": {
            "A": [[values.cell(r, c).value for c in range(5, 9)] for r in range(2, 6)],
            "b": [values.cell(r, 9).value for r in range(2, 6)],
            "A_inverse": [
                [values.cell(r, c).value for c in range(5, 9)] for r in range(8, 12)
            ],
            "solution": [values.cell(r, 9).value for r in range(8, 12)],
        },
        "equilibrium_2": {
            "A": [[values.cell(r, c).value for c in range(5, 9)] for r in range(14, 18)],
            "b": [values.cell(r, 9).value for r in range(14, 18)],
            "A_inverse": [
                [values.cell(r, c).value for c in range(5, 9)] for r in range(20, 24)
            ],
            "solution": [values.cell(r, 9).value for r in range(20, 24)],
        },
        "array_formulas": {
            address: {
                "formula": array_formula_text(formulas[address].value),
                "range": getattr(formulas[address].value, "ref", None),
            }
            for address in ("E8", "I8", "E20", "I20")
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--static", type=Path, required=True)
    parser.add_argument("--dynamic", type=Path, required=True)
    parser.add_argument("--static-slides", type=Path)
    parser.add_argument("--dynamic-slides", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)

    extract_dynamic(args.dynamic).to_csv(
        args.output / "dynamic_excel_reference.csv", index=False
    )
    (args.output / "static_excel_reference.json").write_text(
        json.dumps(extract_static(args.static), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    fingerprints = {
        args.static.name: sha256(args.static),
        args.dynamic.name: sha256(args.dynamic),
    }
    for slide_path in (args.static_slides, args.dynamic_slides):
        if slide_path is not None:
            fingerprints[slide_path.name] = sha256(slide_path)
    (args.output / "source_fingerprints.json").write_text(
        json.dumps(fingerprints, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
